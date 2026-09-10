from datetime import date, datetime
from pathlib import Path

import pytest

from openpyxl import Workbook, load_workbook

from extractor import BatchExtractionError, allocate_quotas, extract_files, parse_date, read_candidates, sha256_file


def make_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "With"
    ws.append(["医疗类别名称", "住院或门诊号", "就诊ID", "人员姓名", "个人编号", "身份证号", "入院日期", "出院日期", "结算日期", "医保目录名称", "医保目录编码", "诊断名称"])
    rows = [
        ["普通门诊", "M1", "V1", "张三", "P1", "ID1", datetime(2024, 1, 1), None, datetime(2024, 1, 2), "项目A", "A", "诊断甲"],
        ["普通门诊", "M1", "V1", "张三", "P1", "ID1", datetime(2024, 1, 1), None, datetime(2024, 1, 2), "项目B", "B", "诊断甲"],
        ["住院", "H1", "V2", "李四", "P2", "ID2", datetime(2025, 6, 1), datetime(2025, 6, 8), datetime(2025, 6, 9), "项目C", "C", "诊断乙"],
        ["门诊慢特病", "M2", "V3", "王五", "P3", "ID3", None, None, datetime(2025, 9, 1), "项目D", "D", "诊断丙"],
        ["急诊", "M3", "V4", "赵六", "P4", "ID4", datetime(2026, 12, 31), None, datetime(2027, 1, 1), "项目E", "E", "诊断丁"],
        ["住院", "H2", "V5", "钱七", "P5", "ID5", datetime(2026, 5, 1), None, None, "项目F", "F", "诊断戊"],
        ["住院", "OLD", "V6", "孙八", "P6", "ID6", datetime(2023, 12, 31), None, None, "项目G", "G", "诊断己"],
    ]
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_quota_rules():
    assert allocate_quotas([2024, 2025, 2026], 5) == {2024: 1, 2025: 2, 2026: 2}
    assert allocate_quotas([2023, 2024, 2025, 2026], 2) == {2023: 1, 2024: 1}


def test_parse_slash_datetime_with_milliseconds():
    parsed = parse_date("2024/05/18 09:47:56.000")
    assert parsed == datetime(2024, 5, 18, 9, 47, 56)


def test_end_to_end_privacy_merge_and_source_unchanged(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    before = sha256_file(source)
    output = extract_files([source], date(2024, 1, 1), date(2026, 12, 31), 5, 12345, tmp_path / "out")
    assert sha256_file(source) == before
    wb = load_workbook(output, data_only=True)
    ws = wb["抽取结果"]
    headers = [cell.value for cell in ws[1]]
    assert "人员姓名" in headers
    assert "诊断名称" in headers
    assert "个人编号" not in headers
    assert "身份证号" not in headers
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 5
    merged = next(row for row in rows if row["就诊ID"] == "V1")
    assert merged["人员姓名"] == "张三"
    assert merged["医保目录名称"] == "项目A；项目B"
    assert merged["医保目录编码"] == "A；B"
    assert {row["归属年份"] for row in rows} == {2024, 2025, 2026}
    assert next(row for row in rows if row["就诊ID"] == "V4")["归属年份"] == 2026
    params = dict(wb["参数"].iter_rows(min_row=2, values_only=True))
    assert params["工具版本"] == "1.6"
    wb.close()


def test_same_seed_same_visits(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    outputs = [extract_files([source], date(2024, 1, 1), date(2026, 12, 31), 3, 77, tmp_path / f"out{i}") for i in range(2)]
    visits = []
    for output in outputs:
        wb = load_workbook(output, data_only=True)
        ws = wb["抽取结果"]
        headers = [cell.value for cell in ws[1]]
        index = headers.index("就诊ID")
        visits.append([row[index] for row in ws.iter_rows(min_row=2, values_only=True)])
        wb.close()
    assert visits[0] == visits[1]


def test_all_files_failed_raises_and_keeps_error_report(tmp_path):
    with pytest.raises(BatchExtractionError) as caught:
        extract_files([tmp_path / "missing.xlsx"], date(2024, 1, 1), date(2026, 12, 31), 5, 1, tmp_path / "out")
    report = caught.value.output_path
    assert report.exists()
    wb = load_workbook(report, data_only=True)
    assert wb["运行汇总"]["F2"].value == "失败"
    assert wb["异常明细"]["A2"].value == "missing.xlsx"
    wb.close()


def test_missing_required_fields_is_rejected(tmp_path):
    source = tmp_path / "invalid.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["入院时间", "姓名"])
    ws.append([datetime(2024, 1, 1), "张三"])
    wb.save(source)
    with pytest.raises(ValueError, match="业务工作表缺少必要字段") as caught:
        read_candidates(source, date(2024, 1, 1), date(2024, 12, 31), 1)
    assert "医疗类别" in str(caught.value)
    assert "医保目录编码" in str(caught.value)


def test_equivalent_date_formats_merge_into_one_record(tmp_path):
    source = tmp_path / "dates.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["医疗类别", "住院号", "就诊ID", "姓名", "入院时间", "出院时间", "医保目录名称", "医保目录编码"])
    ws.append(["住院", "H1", "V1", "张三", datetime(2024, 1, 1), datetime(2024, 1, 2), "项目A", "A"])
    ws.append(["住院", "H1", "V1", "张三", "2024/01/01 00:00:00.000", "2024-01-02 00:00:00", "项目B", "B"])
    wb.save(source)
    records, _, _ = read_candidates(source, date(2024, 1, 1), date(2024, 12, 31), 1)
    assert len(records) == 1
    assert records[0]["医保目录名称"] == "项目A；项目B"


def test_complete_business_sheet_wins_over_earlier_partial_sheet(tmp_path):
    source = tmp_path / "sheets.xlsx"
    wb = Workbook()
    partial = wb.active
    partial.title = "Partial"
    partial.append(["入院时间", "医疗类别", "就诊ID", "姓名", "医保目录名称"])
    complete = wb.create_sheet("Complete")
    complete.append(["入院时间", "医疗类别", "住院号", "就诊ID", "姓名", "医保目录名称", "医保目录编码"])
    complete.append([datetime(2025, 1, 1), "门诊", "M1", "V1", "张三", "项目A", "A"])
    wb.save(source)
    records, _, _ = read_candidates(source, date(2025, 1, 1), date(2025, 12, 31), 1)
    assert len(records) == 1
    assert records[0]["来源工作表"] == "Complete"


def test_item_name_and_code_are_deduplicated_as_pairs(tmp_path):
    source = tmp_path / "items.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["医疗类别", "住院号", "就诊ID", "姓名", "入院时间", "医保目录名称", "医保目录编码"])
    ws.append(["门诊", "M1", "V1", "张三", datetime(2025, 1, 1), "同名项目", "A"])
    ws.append(["门诊", "M1", "V1", "张三", datetime(2025, 1, 1), "同名项目", "B"])
    ws.append(["门诊", "M1", "V1", "张三", datetime(2025, 1, 1), "同名项目", "A"])
    wb.save(source)
    records, _, _ = read_candidates(source, date(2025, 1, 1), date(2025, 12, 31), 1)
    assert records[0]["医保目录名称"] == "同名项目；同名项目"
    assert records[0]["医保目录编码"] == "A；B"


def test_partial_failure_details_can_be_returned(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    output, errors = extract_files(
        [source, tmp_path / "missing.xlsx"],
        date(2024, 1, 1),
        date(2026, 12, 31),
        2,
        1,
        tmp_path / "out",
        return_errors=True,
    )
    assert output.exists()
    assert errors and errors[0][0] == "missing.xlsx"
